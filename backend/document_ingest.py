from __future__ import annotations

import hashlib
import json
import os
from collections.abc import Callable, Iterable
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pypdf import PdfReader

try:
    from pdf2image import convert_from_path
except ImportError:  # pragma: no cover - optional dependency
    convert_from_path = None

try:
    import pytesseract
except ImportError:  # pragma: no cover - optional dependency
    pytesseract = None

try:
    from PIL import Image
except ImportError:  # pragma: no cover - optional dependency
    Image = None

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover - optional dependency
    DocxDocument = None


DEFAULT_DATA_FOLDER = "data"
DEFAULT_MANIFEST_FILE = "pdf_manifest.json"
KNOWN_PDF_FOLDERS = ("Manual_pdfs", "manuals_pdf")
MAX_SOURCE_FILE_BYTES = 25 * 1024 * 1024
MAX_OVERSIZED_SHARD_CHARS = 350_000

PDF_EXTENSIONS = {".pdf"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
DOCX_EXTENSIONS = {".docx"}
TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".rst",
    ".csv",
    ".tsv",
    ".json",
    ".yaml",
    ".yml",
    ".log",
    ".ini",
    ".cfg",
    ".conf",
    ".toml",
    ".xml",
    ".html",
    ".htm",
    ".css",
    ".sql",
    ".bat",
    ".ps1",
    ".py",
    ".js",
    ".ts",
    ".jsx",
    ".tsx",
}
SUPPORTED_EXTENSIONS = PDF_EXTENSIONS | IMAGE_EXTENSIONS | DOCX_EXTENSIONS | TEXT_EXTENSIONS

WINDOWS_TESSERACT_CANDIDATES = (
    Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
    Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
)

WINDOWS_POPPLER_CANDIDATES = (
    Path(r"C:\poppler-25.12.0\Library\bin"),
    Path(r"C:\poppler\Library\bin"),
)


Logger = Callable[[str], None]
ProgressCallback = Callable[[dict[str, Any]], None]
BAD_PDF_QUARANTINE_DIR = "_quarantine_bad_pdfs"


@dataclass(slots=True)
class IngestResult:
    source_path: str
    output_path: str
    status: str
    method: str
    message: str
    source_kind: str = "file"


class InvalidPdfError(RuntimeError):
    pass


def clean_text(text: str) -> str:
    if not text:
        return ""

    lines = [line.rstrip() for line in text.splitlines()]
    return "\n".join([line for line in lines if line.strip()])


def resolve_default_pdf_folder(base_dir: str | Path = ".") -> Path:
    base_path = Path(base_dir)
    for folder_name in KNOWN_PDF_FOLDERS:
        candidate = base_path / folder_name
        if candidate.exists():
            return candidate
    return base_path / KNOWN_PDF_FOLDERS[0]


def supported_extensions() -> tuple[str, ...]:
    return tuple(sorted(SUPPORTED_EXTENSIONS))


def _logger(logger: Logger | None) -> Logger:
    return logger or (lambda _: None)


def _emit_progress(progress_callback: ProgressCallback | None, payload: dict[str, Any]) -> None:
    if not progress_callback:
        return
    try:
        progress_callback(payload)
    except Exception:
        pass


def _canonical_path(path: str | Path) -> str:
    return str(Path(path).resolve())


def _fingerprint(path: str | Path) -> dict[str, float | int]:
    stat = Path(path).stat()
    return {
        "mtime": stat.st_mtime,
        "size": stat.st_size,
    }


def _manifest_entry(
    source_path: Path,
    output_path: Path | None = None,
    output_paths: list[Path] | None = None,
) -> dict[str, Any]:
    payload = _fingerprint(source_path)
    resolved_outputs = output_paths or ([output_path] if output_path is not None else [])
    payload.update(
        {
            "source_name": source_path.name,
            "output_name": resolved_outputs[0].name if resolved_outputs else _output_name_for_source(source_path),
            "output_names": [path.name for path in resolved_outputs],
        }
    )
    return payload


def _normalize_manifest(manifest: dict[str, Any]) -> dict[str, dict[str, Any]]:
    normalized: dict[str, dict[str, Any]] = {}
    for raw_key, raw_value in manifest.items():
        if not isinstance(raw_value, dict):
            continue
        source_path = Path(raw_key)
        entry = {
            "mtime": raw_value.get("mtime"),
            "size": raw_value.get("size"),
            "source_name": raw_value.get("source_name") or source_path.name,
            "output_name": raw_value.get("output_name") or _output_name_for_source(source_path),
            "output_names": raw_value.get("output_names") or [],
        }
        if not entry["output_names"] and entry["output_name"]:
            entry["output_names"] = [entry["output_name"]]
        normalized[_canonical_path(source_path)] = entry
    return normalized


def _entry_matches_fingerprint(entry: dict[str, Any] | None, fingerprint: dict[str, float | int]) -> bool:
    if not entry:
        return False
    return entry.get("mtime") == fingerprint.get("mtime") and entry.get("size") == fingerprint.get("size")


def _has_reusable_output(source_path: Path, output_path: Path) -> bool:
    if not output_path.exists():
        return False
    try:
        output_stat = output_path.stat()
        source_stat = source_path.stat()
    except OSError:
        return False
    return output_stat.st_size > 0 and output_stat.st_mtime >= source_stat.st_mtime


def _entry_output_paths(entry: dict[str, Any], output_dir: Path, source_path: Path | None = None) -> list[Path]:
    output_names = entry.get("output_names") or []
    if not output_names:
        output_name = str(entry.get("output_name") or "")
        if output_name:
            output_names = [output_name]
    if not output_names and source_path is not None:
        output_names = [_output_name_for_source(source_path)]
    return [output_dir / str(name) for name in output_names if str(name).strip()]


def _entry_outputs_ready(source_path: Path, entry: dict[str, Any] | None, output_dir: Path) -> bool:
    if not entry:
        return False
    try:
        source_stat = source_path.stat()
    except OSError:
        return False
    output_paths = _entry_output_paths(entry, output_dir, source_path=source_path)
    if not output_paths:
        return False
    for output_path in output_paths:
        try:
            output_stat = output_path.stat()
        except OSError:
            return False
        if output_stat.st_size <= 0 or output_stat.st_mtime < source_stat.st_mtime:
            return False
    return True


def _find_history_match(
    source_path: Path,
    manifest: dict[str, dict[str, Any]],
    fingerprint: dict[str, float | int],
    output_dir: Path,
) -> dict[str, Any] | None:
    source_name = source_path.name.lower()
    for entry in manifest.values():
        entry_name = str(entry.get("source_name") or "").lower()
        output_paths = _entry_output_paths(entry, output_dir, source_path=source_path)
        if entry_name != source_name or not output_paths:
            continue
        if not _entry_matches_fingerprint(entry, fingerprint):
            continue
        if all(path.exists() for path in output_paths):
            return entry
    return None


def _load_manifest(manifest_path: str | Path) -> dict[str, dict[str, Any]]:
    manifest_file = Path(manifest_path)
    if not manifest_file.exists():
        return {}

    try:
        with manifest_file.open("r", encoding="utf-8") as handle:
            return _normalize_manifest(json.load(handle))
    except Exception:
        return {}


def _save_manifest(manifest: dict[str, dict[str, Any]], manifest_path: str | Path) -> None:
    manifest_file = Path(manifest_path)
    with manifest_file.open("w", encoding="utf-8") as handle:
        json.dump(manifest, handle, indent=2)


def _find_tesseract_cmd() -> str | None:
    env_value = os.environ.get("TESSERACT_CMD")
    if env_value and Path(env_value).exists():
        return env_value

    for candidate in WINDOWS_TESSERACT_CANDIDATES:
        if candidate.exists():
            return str(candidate)

    return None


def _find_poppler_path() -> str | None:
    env_value = os.environ.get("POPPLER_PATH")
    if env_value and Path(env_value).exists():
        return env_value

    for candidate in WINDOWS_POPPLER_CANDIDATES:
        if candidate.exists():
            return str(candidate)

    return None


def source_environment() -> dict[str, str | bool | None]:
    tesseract_cmd = _find_tesseract_cmd()
    poppler_path = _find_poppler_path()

    return {
        "ocr_ready": bool(convert_from_path and pytesseract and tesseract_cmd),
        "tesseract_cmd": tesseract_cmd,
        "poppler_path": poppler_path,
        "pdf2image_installed": bool(convert_from_path),
        "pytesseract_installed": bool(pytesseract),
        "pillow_installed": bool(Image),
        "docx_installed": bool(DocxDocument),
    }


def ocr_environment() -> dict[str, str | bool | None]:
    return source_environment()


def _validate_pdf_file(pdf_path: str | Path) -> None:
    path = Path(pdf_path)
    try:
        with path.open("rb") as handle:
            head = handle.read(8192)
            try:
                handle.seek(-8192, os.SEEK_END)
            except OSError:
                handle.seek(0)
            tail = handle.read(8192)
    except OSError as exc:
        raise RuntimeError(f"Could not read PDF file: {exc}") from exc

    pdf_offset = head.find(b"%PDF-")
    if pdf_offset == -1:
        snippet = head[:32].lstrip().lower()
        if snippet.startswith((b"<html", b"<!doctype", b"<div", b"<body")):
            raise InvalidPdfError(
                f"{path.name} is not a valid PDF file. It appears to contain HTML instead of PDF data."
            )
        raise InvalidPdfError(f"{path.name} is not a valid PDF file. Missing %PDF header.")

    if pdf_offset > 0:
        raise InvalidPdfError(
            f"{path.name} is corrupted. Found non-PDF content before the %PDF header at byte {pdf_offset}."
        )

    if b"%%EOF" not in tail:
        raise InvalidPdfError(f"{path.name} is truncated or corrupted. Missing %%EOF marker.")


def _quarantine_destination(source_path: Path) -> Path:
    quarantine_dir = source_path.parent / BAD_PDF_QUARANTINE_DIR
    quarantine_dir.mkdir(parents=True, exist_ok=True)
    candidate = quarantine_dir / source_path.name
    if not candidate.exists():
        return candidate

    stem = source_path.stem
    suffix = source_path.suffix
    counter = 1
    while True:
        candidate = quarantine_dir / f"{stem}_{counter}{suffix}"
        if not candidate.exists():
            return candidate
        counter += 1


def _quarantine_invalid_pdf(source_path: Path, reason: str, logger: Logger | None = None) -> Path:
    destination = _quarantine_destination(source_path)
    source_path.replace(destination)
    _logger(logger)(f"Quarantined bad PDF -> {destination.name} ({reason})")
    return destination


def _extract_pdf_native(pdf_path: str | Path, logger: Logger | None = None) -> str:
    log = _logger(logger)
    _validate_pdf_file(pdf_path)
    reader = PdfReader(str(pdf_path))
    pages: list[str] = []

    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"\n--- Page {index} ---\n{text}")
            else:
                log(f"No native text on page {index} of {Path(pdf_path).name}")
        except Exception as exc:
            log(f"Native extraction failed on page {index} of {Path(pdf_path).name}: {exc}")

    return clean_text("\n".join(pages))


def _extract_pdf_ocr(pdf_path: str | Path, logger: Logger | None = None) -> str:
    log = _logger(logger)
    _validate_pdf_file(pdf_path)
    env = source_environment()

    if not env["pdf2image_installed"]:
        raise RuntimeError("pdf2image is not installed, so PDF OCR cannot run.")
    if not env["pytesseract_installed"]:
        raise RuntimeError("pytesseract is not installed, so PDF OCR cannot run.")
    if not env["tesseract_cmd"]:
        raise RuntimeError("Tesseract was not found. Set TESSERACT_CMD or install Tesseract-OCR.")

    pytesseract.pytesseract.tesseract_cmd = str(env["tesseract_cmd"])

    convert_kwargs = {}
    if env["poppler_path"]:
        convert_kwargs["poppler_path"] = env["poppler_path"]

    images = convert_from_path(str(pdf_path), **convert_kwargs)
    pages: list[str] = []

    for index, image in enumerate(images, start=1):
        try:
            text = pytesseract.image_to_string(image) or ""
            if text.strip():
                pages.append(f"\n--- OCR Page {index} ---\n{text}")
            else:
                log(f"OCR found no text on page {index} of {Path(pdf_path).name}")
        except Exception as exc:
            log(f"OCR failed on page {index} of {Path(pdf_path).name}: {exc}")

    return clean_text("\n".join(pages))


def _extract_pdf_text(
    pdf_path: str | Path,
    method: str = "auto",
    logger: Logger | None = None,
) -> tuple[str, str]:
    log = _logger(logger)
    selected_method = method.lower().strip()

    if selected_method not in {"auto", "native", "ocr"}:
        raise ValueError("method must be 'auto', 'native', or 'ocr'")

    if selected_method in {"auto", "native"}:
        native_text = _extract_pdf_native(pdf_path, logger=logger)
        if native_text.strip():
            return native_text, "native"
        if selected_method == "native":
            return "", "native"
        log(f"Falling back to OCR for {Path(pdf_path).name}")

    return _extract_pdf_ocr(pdf_path, logger=logger), "ocr"


def _extract_image_text(image_path: str | Path) -> tuple[str, str]:
    env = source_environment()

    if not env["pytesseract_installed"]:
        raise RuntimeError("pytesseract is not installed, so image OCR cannot run.")
    if not env["pillow_installed"]:
        raise RuntimeError("Pillow is not installed, so image OCR cannot run.")
    if not env["tesseract_cmd"]:
        raise RuntimeError("Tesseract was not found. Set TESSERACT_CMD or install Tesseract-OCR.")

    pytesseract.pytesseract.tesseract_cmd = str(env["tesseract_cmd"])
    with Image.open(str(image_path)) as image:
        text = pytesseract.image_to_string(image) or ""
    return clean_text(text), "image-ocr"


def _read_plain_text(path: str | Path) -> tuple[str, str]:
    attempts = ("utf-8", "utf-8-sig", "latin-1")
    last_error = None

    for encoding in attempts:
        try:
            with Path(path).open("r", encoding=encoding) as handle:
                return clean_text(handle.read()), f"text:{encoding}"
        except UnicodeDecodeError as exc:
            last_error = exc

    raise RuntimeError(f"Unable to decode text file: {last_error}")


def _extract_docx_text(path: str | Path) -> tuple[str, str]:
    if not DocxDocument:
        raise RuntimeError("python-docx is not installed, so DOCX import cannot run.")

    document = DocxDocument(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return clean_text("\n".join(parts)), "docx"


def _output_name_for_source(source_path: Path) -> str:
    suffix = source_path.suffix.lower()
    if suffix == ".pdf":
        return f"{source_path.stem}.txt"

    tag = suffix.lstrip(".") or "file"
    safe_stem = source_path.stem.strip() or "source"
    digest = hashlib.md5(str(source_path.resolve()).encode("utf-8")).hexdigest()[:8]
    return f"{safe_stem}__{tag}__{digest}.txt"


def _output_name_for_source_part(source_path: Path, part_index: int) -> str:
    base_name = _output_name_for_source(source_path)
    stem = Path(base_name).stem
    return f"{stem}__part{part_index:04d}.txt"


def _result(
    source_path: Path,
    output_path: Path,
    status: str,
    method: str,
    message: str,
    source_kind: str,
) -> IngestResult:
    return IngestResult(
        source_path=str(source_path),
        output_path=str(output_path),
        status=status,
        method=method,
        message=message,
        source_kind=source_kind,
    )


def _extract_supported_text(
    source_path: str | Path,
    method: str = "auto",
    logger: Logger | None = None,
) -> tuple[str, str, str]:
    path = Path(source_path)
    suffix = path.suffix.lower()

    if suffix in PDF_EXTENSIONS:
        text, used_method = _extract_pdf_text(path, method=method, logger=logger)
        return text, used_method, "pdf"
    if suffix in IMAGE_EXTENSIONS:
        text, used_method = _extract_image_text(path)
        return text, used_method, "image"
    if suffix in DOCX_EXTENSIONS:
        text, used_method = _extract_docx_text(path)
        return text, used_method, "docx"
    if suffix in TEXT_EXTENSIONS:
        text, used_method = _read_plain_text(path)
        return text, used_method, "text"

    raise ValueError(f"Unsupported file type: {path.suffix or 'no extension'}")


def _split_text_into_shards(text: str, max_chars: int = MAX_OVERSIZED_SHARD_CHARS) -> list[str]:
    cleaned = clean_text(text)
    if not cleaned:
        return []
    if len(cleaned) <= max_chars:
        return [cleaned]

    paragraphs = [part.strip() for part in cleaned.split("\n\n") if part.strip()]
    if not paragraphs:
        return [cleaned]

    shards: list[str] = []
    current: list[str] = []
    current_length = 0

    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                shards.append("\n\n".join(current).strip())
                current = []
                current_length = 0
            start = 0
            while start < len(paragraph):
                chunk = paragraph[start : start + max_chars].strip()
                if chunk:
                    shards.append(chunk)
                start += max_chars
            continue

        projected = current_length + len(paragraph) + (2 if current else 0)
        if current and projected > max_chars:
            shards.append("\n\n".join(current).strip())
            current = [paragraph]
            current_length = len(paragraph)
        else:
            current.append(paragraph)
            current_length = projected

    if current:
        shards.append("\n\n".join(current).strip())
    return [shard for shard in shards if shard]


def _remove_entry_outputs(entry: dict[str, Any] | None, output_dir: Path) -> None:
    if not entry:
        return
    for output_path in _entry_output_paths(entry, output_dir):
        try:
            if output_path.exists():
                output_path.unlink()
        except OSError:
            pass


def _process_source_file(
    source_path: Path,
    output_path: Path,
    *,
    method: str,
    logger: Logger | None,
) -> tuple[IngestResult, dict[str, Any] | None]:
    try:
        text, used_method, source_kind = _extract_supported_text(source_path, method=method, logger=logger)
        if not text.strip():
            message = f"No text extracted from {source_path.name}"
            _logger(logger)(message)
            return _result(source_path, output_path, "failed", used_method, message, source_kind), None

        output_path.write_text(text, encoding="utf-8")
        message = f"Saved TXT: {output_path.name}"
        _logger(logger)(message)
        return _result(source_path, output_path, "converted", used_method, message, source_kind), _manifest_entry(
            source_path, output_path
        )
    except Exception as exc:
        message = f"Failed to process {source_path.name}: {exc}"
        _logger(logger)(message)
        return _result(source_path, output_path, "failed", method, message, source_path.suffix.lower()), None


def _process_oversized_source_file(
    source_path: Path,
    output_dir: Path,
    *,
    method: str,
    logger: Logger | None,
    previous_entry: dict[str, Any] | None = None,
) -> tuple[list[IngestResult], dict[str, Any] | None]:
    log = _logger(logger)
    try:
        text, used_method, source_kind = _extract_supported_text(source_path, method=method, logger=logger)
        if not text.strip():
            output_path = output_dir / _output_name_for_source(source_path)
            message = f"No text extracted from oversized source {source_path.name}"
            log(message)
            return [_result(source_path, output_path, "failed", used_method, message, source_kind)], None

        shards = _split_text_into_shards(text)
        if not shards:
            output_path = output_dir / _output_name_for_source(source_path)
            message = f"No usable text shards for oversized source {source_path.name}"
            log(message)
            return [_result(source_path, output_path, "failed", used_method, message, source_kind)], None

        _remove_entry_outputs(previous_entry, output_dir)

        output_paths: list[Path] = []
        results: list[IngestResult] = []
        for index, shard in enumerate(shards, start=1):
            output_path = output_dir / _output_name_for_source_part(source_path, index)
            output_path.write_text(shard, encoding="utf-8")
            output_paths.append(output_path)
            message = f"Saved shard {index}/{len(shards)}: {output_path.name}"
            log(message)
            results.append(_result(source_path, output_path, "converted", used_method, message, source_kind))

        summary_path = output_paths[0]
        summary_message = f"Split oversized source into {len(output_paths)} shard(s): {source_path.name}"
        log(summary_message)
        results.append(_result(source_path, summary_path, "split", used_method, summary_message, source_kind))
        return results, _manifest_entry(source_path, output_paths=output_paths)
    except Exception as exc:
        output_path = output_dir / _output_name_for_source(source_path)
        message = f"Failed to process oversized source {source_path.name}: {exc}"
        log(message)
        return [_result(source_path, output_path, "failed", method, message, source_path.suffix.lower())], None


def _max_parallel_workers(method: str, pending_count: int) -> int:
    if pending_count <= 1:
        return 1
    cpu_count = os.cpu_count() or 1
    if method == "ocr":
        return 1
    if method == "auto":
        return max(1, min(2, pending_count, cpu_count))
    return max(1, min(4, pending_count, cpu_count))


def _is_source_too_large(source_path: Path) -> bool:
    try:
        return source_path.stat().st_size > MAX_SOURCE_FILE_BYTES
    except OSError:
        return False


def collect_supported_files(folder: str | Path, recursive: bool = True) -> list[Path]:
    base_folder = Path(folder)
    if not base_folder.exists():
        return []

    iterator = base_folder.rglob("*") if recursive else base_folder.glob("*")
    files = [path.resolve() for path in iterator if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS]
    return sorted(set(files))


def ingest_source_files(
    source_files: Iterable[str | Path],
    output_folder: str | Path = DEFAULT_DATA_FOLDER,
    manifest_path: str | Path = DEFAULT_MANIFEST_FILE,
    method: str = "auto",
    logger: Logger | None = None,
    delete_stale: bool = False,
    progress_callback: ProgressCallback | None = None,
) -> list[IngestResult]:
    log = _logger(logger)
    output_dir = Path(output_folder)
    output_dir.mkdir(parents=True, exist_ok=True)

    manifest = _load_manifest(manifest_path)
    results: list[IngestResult] = []
    normalized_inputs: list[Path] = []

    for raw_path in source_files:
        source_path = Path(raw_path)
        if source_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        if source_path.exists():
            normalized_inputs.append(source_path.resolve())
        else:
            output_path = output_dir / _output_name_for_source(source_path)
            message = f"Missing source file: {source_path}"
            log(message)
            results.append(_result(source_path, output_path, "failed", "none", message, "missing"))

    normalized_inputs = sorted(set(normalized_inputs))
    seen_sources = {_canonical_path(path) for path in normalized_inputs}
    pending_unchanged_logs: list[str] = []
    total_sources = len(normalized_inputs)
    completed_sources = 0
    manifest_dirty = False
    pending_sources: list[tuple[Path, Path]] = []

    fingerprint_index: dict[tuple[str, float | int, float | int], dict[str, Any]] = {}
    for entry in manifest.values():
        source_name = str(entry.get("source_name") or "").lower()
        mtime = entry.get("mtime")
        size = entry.get("size")
        if source_name and mtime is not None and size is not None:
            fingerprint_index[(source_name, mtime, size)] = entry

    def mark_progress(source_path: Path, status: str, message: str) -> None:
        nonlocal completed_sources
        completed_sources += 1
        _emit_progress(
            progress_callback,
            {
                "current": completed_sources,
                "total": total_sources,
                "source_name": source_path.name,
                "status": status,
                "message": message,
            },
        )

    for source_path in normalized_inputs:
        source_key = _canonical_path(source_path)
        output_path = output_dir / _output_name_for_source(source_path)
        current_fingerprint = _fingerprint(source_path)
        existing_entry = manifest.get(source_key)
        current_entry = _manifest_entry(source_path, output_path)

        if source_path.suffix.lower() in PDF_EXTENSIONS:
            try:
                _validate_pdf_file(source_path)
            except InvalidPdfError as exc:
                try:
                    _remove_entry_outputs(existing_entry, output_dir)
                    if source_key in manifest:
                        manifest.pop(source_key, None)
                        manifest_dirty = True
                    quarantined_path = _quarantine_invalid_pdf(source_path, str(exc), logger=logger)
                    message = (
                        f"Quarantined bad PDF {source_path.name} -> {quarantined_path.parent.name}\\{quarantined_path.name}. "
                        f"{exc}"
                    )
                    results.append(_result(source_path, quarantined_path, "quarantined", "quarantine", message, "pdf"))
                    mark_progress(source_path, "quarantined", message)
                except OSError as move_exc:
                    message = (
                        f"Detected bad PDF {source_path.name}, but failed to quarantine it: {move_exc}. "
                        f"Reason: {exc}"
                    )
                    _logger(logger)(message)
                    results.append(_result(source_path, output_path, "failed", "quarantine", message, "pdf"))
                    mark_progress(source_path, "failed", message)
                continue

        if _entry_matches_fingerprint(existing_entry, current_fingerprint) and _entry_outputs_ready(
            source_path, existing_entry, output_dir
        ):
            message = f"No change: {source_path.name}"
            pending_unchanged_logs.append(message)
            results.append(
                _result(source_path, output_path, "unchanged", "cached", message, source_path.suffix.lower())
            )
            mark_progress(source_path, "unchanged", message)
            continue

        history_match = fingerprint_index.get(
            (source_path.name.lower(), current_fingerprint["mtime"], current_fingerprint["size"])
        )
        if history_match and not (output_dir / str(history_match.get("output_name") or "")).exists():
            history_match = None
        if history_match is None:
            history_match = _find_history_match(source_path, manifest, current_fingerprint, output_dir)
        if history_match:
            manifest[source_key] = current_entry
            manifest_dirty = True
            message = f"No change: {source_path.name}"
            pending_unchanged_logs.append(message)
            results.append(
                _result(source_path, output_path, "unchanged", "cached-history", message, source_path.suffix.lower())
            )
            mark_progress(source_path, "unchanged", message)
            continue

        if _has_reusable_output(source_path, output_path):
            manifest[source_key] = current_entry
            manifest_dirty = True
            message = f"No change: {source_path.name}"
            pending_unchanged_logs.append(message)
            results.append(
                _result(source_path, output_path, "unchanged", "cached-output", message, source_path.suffix.lower())
            )
            mark_progress(source_path, "unchanged", message)
            continue

        if _is_source_too_large(source_path):
            size_mb = source_path.stat().st_size / (1024 * 1024)
            message = (
                f"Processing oversized source by splitting: {source_path.name} "
                f"({size_mb:.1f} MB > {MAX_SOURCE_FILE_BYTES / (1024 * 1024):.0f} MB limit)"
            )
            log(message)
            oversized_results, manifest_entry = _process_oversized_source_file(
                source_path,
                output_dir,
                method=method,
                logger=logger,
                previous_entry=existing_entry,
            )
            results.extend(oversized_results)
            if manifest_entry is not None:
                manifest[source_key] = manifest_entry
                manifest_dirty = True
            mark_progress(source_path, oversized_results[-1].status, oversized_results[-1].message)
            continue

        log(f"Processing source: {source_path.name}")
        pending_sources.append((source_path, output_path))

    max_workers = _max_parallel_workers(method.lower().strip(), len(pending_sources))
    if max_workers == 1:
        for source_path, output_path in pending_sources:
            result, manifest_entry = _process_source_file(
                source_path,
                output_path,
                method=method,
                logger=logger,
            )
            results.append(result)
            if manifest_entry is not None:
                manifest[_canonical_path(source_path)] = manifest_entry
                manifest_dirty = True
            mark_progress(source_path, result.status, result.message)
    else:
        with ThreadPoolExecutor(max_workers=max_workers, thread_name_prefix="pdf-ingest") as executor:
            futures = {
                executor.submit(
                    _process_source_file,
                    source_path,
                    output_path,
                    method=method,
                    logger=logger,
                ): (source_path, output_path)
                for source_path, output_path in pending_sources
            }
            for future in as_completed(futures):
                source_path, _output_path = futures[future]
                result, manifest_entry = future.result()
                results.append(result)
                if manifest_entry is not None:
                    manifest[_canonical_path(source_path)] = manifest_entry
                    manifest_dirty = True
                mark_progress(source_path, result.status, result.message)

    if pending_unchanged_logs:
        sample = ", ".join(Path(message.replace("No change: ", "")).name for message in pending_unchanged_logs[:3])
        remaining = len(pending_unchanged_logs) - min(len(pending_unchanged_logs), 3)
        if remaining > 0:
            log(f"Skipped {len(pending_unchanged_logs)} previously processed source(s): {sample}, +{remaining} more")
        else:
            log(f"Skipped {len(pending_unchanged_logs)} previously processed source(s): {sample}")

    if delete_stale:
        stale_sources = [source for source in manifest.keys() if source not in seen_sources]
        for source in stale_sources:
            stale_entry = manifest.get(source)
            removed_any = False
            for output_path in _entry_output_paths(stale_entry or {}, output_dir, source_path=Path(source)):
                if output_path.exists():
                    try:
                        output_path.unlink()
                        removed_any = True
                        message = f"Removed stale TXT: {output_path.name}"
                        log(message)
                        results.append(
                            IngestResult(
                                source_path=source,
                                output_path=str(output_path),
                                status="removed",
                                method="cleanup",
                                message=message,
                                source_kind="cleanup",
                            )
                        )
                    except Exception as exc:
                        log(f"Failed removing stale TXT for {source}: {exc}")
            if not removed_any and stale_entry:
                legacy_output_path = output_dir / _output_name_for_source(Path(source))
                if legacy_output_path.exists():
                    try:
                        legacy_output_path.unlink()
                    except Exception as exc:
                        log(f"Failed removing stale TXT for {source}: {exc}")
            manifest.pop(source, None)
            manifest_dirty = True

    if manifest_dirty or delete_stale:
        _save_manifest(manifest, manifest_path)
    return results


def ingest_source_folder(
    source_folder: str | Path,
    output_folder: str | Path = DEFAULT_DATA_FOLDER,
    manifest_path: str | Path = DEFAULT_MANIFEST_FILE,
    method: str = "auto",
    logger: Logger | None = None,
    recursive: bool = True,
    delete_stale: bool = False,
    progress_callback: ProgressCallback | None = None,
) -> list[IngestResult]:
    source_files = collect_supported_files(source_folder, recursive=recursive)
    return ingest_source_files(
        source_files=source_files,
        output_folder=output_folder,
        manifest_path=manifest_path,
        method=method,
        logger=logger,
        delete_stale=delete_stale,
        progress_callback=progress_callback,
    )


def ingest_pdf_files(
    pdf_files: Iterable[str | Path],
    output_folder: str | Path = DEFAULT_DATA_FOLDER,
    manifest_path: str | Path = DEFAULT_MANIFEST_FILE,
    method: str = "auto",
    logger: Logger | None = None,
    delete_stale: bool = False,
    progress_callback: ProgressCallback | None = None,
) -> list[IngestResult]:
    pdf_only = [path for path in pdf_files if Path(path).suffix.lower() in PDF_EXTENSIONS]
    return ingest_source_files(
        source_files=pdf_only,
        output_folder=output_folder,
        manifest_path=manifest_path,
        method=method,
        logger=logger,
        delete_stale=delete_stale,
        progress_callback=progress_callback,
    )


def ingest_pdf_folder(
    pdf_folder: str | Path | None = None,
    output_folder: str | Path = DEFAULT_DATA_FOLDER,
    manifest_path: str | Path = DEFAULT_MANIFEST_FILE,
    method: str = "auto",
    logger: Logger | None = None,
    delete_stale: bool = True,
    progress_callback: ProgressCallback | None = None,
) -> list[IngestResult]:
    base_folder = Path(pdf_folder) if pdf_folder else resolve_default_pdf_folder()
    base_folder.mkdir(parents=True, exist_ok=True)
    pdf_files = sorted(base_folder.glob("*.pdf"))
    return ingest_pdf_files(
        pdf_files=pdf_files,
        output_folder=output_folder,
        manifest_path=manifest_path,
        method=method,
        logger=logger,
        delete_stale=delete_stale,
        progress_callback=progress_callback,
    )


def summarize_ingest(results: Iterable[IngestResult]) -> dict[str, int]:
    summary = {
        "converted": 0,
        "unchanged": 0,
        "failed": 0,
        "removed": 0,
        "quarantined": 0,
    }

    for result in results:
        summary[result.status] = summary.get(result.status, 0) + 1

    return summary
